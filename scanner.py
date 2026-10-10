"""聊天记录关键词扫描核心。

支持两类工作表（自动按列名识别，识别不到时按位置回退）：
  1) 含「消息内容」列的工作表（业务部门）
  2) 含「聊天记录」列的工作表（贵金属）
"""
import re
from dataclasses import dataclass, field
from typing import Callable, List, Optional

from openpyxl import load_workbook

# 关键词分隔符：顿号、逗号、分号、竖线、换行、制表符
SPLIT_RE = re.compile(r"[、,，;；|\n\r\t]+")
DATE_RE = re.compile(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}")

CONTENT_HEADERS = ("消息内容", "聊天记录")
NAME_HEADERS = ("行员姓名", "员工姓名", "姓名")
CUSTOMER_HEADERS = ("客户姓名",)
TIME_HEADERS = ("消息时间", "日期", "时间")

# Excel 单元格字符上限，达到该长度说明原文很可能被截断
EXCEL_CELL_LIMIT = 32767
TRUNCATE_THRESHOLD = 32700


@dataclass
class Target:
    sheet: str
    index: int
    col: int          # 0-based 内容列
    mode: str         # 消息内容 / 聊天记录
    headers: List[str] = field(default_factory=list)
    name_col: Optional[int] = None
    customer_col: Optional[int] = None
    time_col: Optional[int] = None
    rows: int = 0     # 有效数据行数


@dataclass
class HitRecord:
    sheet: str
    row: int
    person: str
    peer: str
    date: str
    keywords: List[str]
    counts: dict
    context: str
    text: str
    truncated: bool = False


@dataclass
class ScanResult:
    hits: List[HitRecord] = field(default_factory=list)
    scanned: int = 0
    targets: List[Target] = field(default_factory=list)
    truncated_rows: int = 0
    cancelled: bool = False


def parse_keywords(raw: str) -> List[str]:
    """把「A、B、C」解析为关键词列表，去重并保持顺序。"""
    seen, out = set(), []
    for kw in SPLIT_RE.split(raw or ""):
        kw = kw.strip()
        if kw and kw not in seen:
            seen.add(kw)
            out.append(kw)
    return out


def _norm(v) -> str:
    return str(v).strip() if v is not None else ""


def _find_col(headers: List[str], names) -> Optional[int]:
    for name in names:
        if name in headers:
            return headers.index(name)
    return None


def _detect_targets(wb) -> List[Target]:
    targets = []
    for idx, ws in enumerate(wb.worksheets):
        raw = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()) or ()
        headers = [_norm(h) for h in raw]
        col, mode = None, None
        for name in CONTENT_HEADERS:
            if name in headers:
                col, mode = headers.index(name), name
                break
        if col is None:
            # 回退：第 1 个工作表取 F 列(消息内容)，第 2 个工作表取 B 列(聊天记录)
            if idx == 0 and len(headers) >= 6:
                col, mode = 5, "消息内容"
            elif idx == 1 and len(headers) >= 2:
                col, mode = 1, "聊天记录"
        if col is None:
            continue
        targets.append(
            Target(
                sheet=ws.title,
                index=idx,
                col=col,
                mode=mode,
                headers=headers,
                name_col=_find_col(headers, NAME_HEADERS),
                customer_col=_find_col(headers, CUSTOMER_HEADERS),
                time_col=_find_col(headers, TIME_HEADERS),
            )
        )
    return targets


def _count_rows(wb, targets: List[Target]) -> None:
    """第一遍：统计有效行数，用于进度条。"""
    for t in targets:
        ws = wb.worksheets[t.index]
        n = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if t.col < len(row) and row[t.col] not in (None, ""):
                n += 1
        t.rows = n


def _peer_and_date(text: str, mode: str, person: str, row, t: Target):
    if mode == "消息内容":
        peer = _norm(row[t.customer_col]) if (t.customer_col is not None and t.customer_col < len(row)) else ""
        date = _norm(row[t.time_col]) if (t.time_col is not None and t.time_col < len(row)) else ""
        if not date:
            m = DATE_RE.search(text)
            date = m.group(0) if m else ""
        return peer, date
    # 聊天记录：首行形如 "阚朝阳(南京银行) 和 FX02(交易中心)"
    first = text.split("\n", 1)[0].strip()
    peer = first
    if " 和 " in first:
        parts = [p.strip() for p in first.split(" 和 ")]
        other = [p for p in parts if not (person and person in p)]
        peer = (other[0] if other else parts[-1])
    m = DATE_RE.search(text)
    return peer, (m.group(0) if m else "")


def scan_file(
    path: str,
    keywords: List[str],
    ignore_case: bool = False,
    progress: Optional[Callable[[int, int, int], None]] = None,
    should_cancel: Optional[Callable[[], bool]] = None,
) -> ScanResult:
    """扫描 Excel，返回命中记录。progress(已处理, 总数, 已命中) 由调用方在 UI 线程消费。"""
    result = ScanResult()

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        result.targets = _detect_targets(wb)
        _count_rows(wb, result.targets)
    finally:
        wb.close()

    total = sum(t.rows for t in result.targets)
    if total == 0:
        return result

    needles = [(k, k.lower() if ignore_case else k) for k in keywords]
    needle_map = dict(needles)

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        processed = 0
        for t in result.targets:
            ws = wb.worksheets[t.index]
            last_person = ""
            for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if t.col >= len(row):
                    continue
                cell = row[t.col]
                if cell is None or cell == "":
                    continue
                processed += 1
                text = str(cell)
                truncated = len(text) >= TRUNCATE_THRESHOLD

                if t.name_col is not None and t.name_col < len(row):
                    p = _norm(row[t.name_col])
                    if p:
                        last_person = p
                person = last_person

                hay = text.lower() if ignore_case else text
                kws, counts = [], {}
                for kw, nd in needles:
                    if nd in hay:
                        kws.append(kw)
                        counts[kw] = hay.count(nd)
                if kws:
                    pos = hay.find(needle_map[kws[0]])
                    start = max(0, pos - 40)
                    context = text[start: pos + len(kws[0]) + 40].replace("\n", " ⏎ ")
                    peer, date = _peer_and_date(text, t.mode, person, row, t)
                    result.hits.append(
                        HitRecord(
                            sheet=t.sheet,
                            row=r_idx,
                            person=person,
                            peer=peer,
                            date=date,
                            keywords=kws,
                            counts=counts,
                            context=context,
                            text=text,
                            truncated=truncated,
                        )
                    )
                    if truncated:
                        result.truncated_rows += 1

                if progress and (processed % 100 == 0 or processed == total):
                    progress(processed, total, len(result.hits))
                if should_cancel and should_cancel():
                    result.cancelled = True
                    break
            if result.cancelled:
                break
        result.scanned = processed
    finally:
        wb.close()

    return result
