"""聊天记录合规关键词筛查工具（tkinter 图形界面）。"""
import os
import queue
import threading
import time
import tkinter as tk
from collections import Counter
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

import config
from exporter import export_hits_only, export_marked
from merge_ui import MergeTab
from scanner import HitRecord, parse_keywords, scan_file

BG_HIT = "#FFF3A3"


class App:
    def __init__(self, root: tk.Tk, container=None):
        self.root = root
        self.container = container or root  # 界面放进「筛查」页签
        root.title("聊天记录合规筛查工具")

        self.file_path = ""
        self.result = None
        self.scanning = False
        self.cancel_event = threading.Event()
        self.q = queue.Queue()
        self.case_sensitive = tk.BooleanVar(value=False)

        self._build_ui()
        saved = config.load_keywords()
        if saved:
            self.kw_text.insert("1.0", saved)
        self._update_kw_count()
        root.protocol("WM_DELETE_WINDOW", self.on_close)

    # ---------------- UI ----------------
    def _build_ui(self):
        pad = {"padx": 8, "pady": 5}

        # 1. 文件
        f1 = ttk.LabelFrame(self.container, text="① 待检查的 Excel 文件")
        f1.pack(fill="x", **pad)
        self.file_var = tk.StringVar()
        ttk.Entry(f1, textvariable=self.file_var).pack(side="left", fill="x", expand=True, padx=(8, 4), pady=8)
        ttk.Button(f1, text="浏览…", command=self.choose_file).pack(side="left", padx=4)
        self.file_info = ttk.Label(f1, text="未选择文件", foreground="#666666")
        self.file_info.pack(side="left", padx=8)

        # 2. 关键词
        f2 = ttk.LabelFrame(self.container, text="② 关键词（用「、」分隔，也支持换行 / 逗号 / 分号）")
        f2.pack(fill="both", expand=False, **pad)
        wrap = ttk.Frame(f2)
        wrap.pack(fill="both", expand=True, padx=8, pady=(8, 4))
        self.kw_text = tk.Text(wrap, height=6, wrap="word", undo=True)
        self.kw_text.pack(side="left", fill="both", expand=True)
        sb = ttk.Scrollbar(wrap, command=self.kw_text.yview)
        sb.pack(side="left", fill="y")
        self.kw_text.configure(yscrollcommand=sb.set)
        self.kw_text.bind("<KeyRelease>", lambda e: self._update_kw_count())

        bar = ttk.Frame(f2)
        bar.pack(fill="x", padx=8, pady=(0, 8))
        self.kw_count = ttk.Label(bar, text="0 个关键词", foreground="#666666")
        self.kw_count.pack(side="left")
        ttk.Button(bar, text="从 txt 导入", command=self.import_txt).pack(side="left", padx=8)
        ttk.Button(bar, text="清空", command=self.clear_kw).pack(side="left")
        ttk.Checkbutton(bar, text="区分大小写", variable=self.case_sensitive).pack(side="left", padx=12)
        ttk.Label(bar, text="（关键词会自动保存在本机，下次打开自动载入）", foreground="#888888").pack(side="left", padx=8)

        # 3. 执行
        f3 = ttk.LabelFrame(self.container, text="③ 开始筛查")
        f3.pack(fill="x", **pad)
        self.btn_scan = ttk.Button(f3, text="开始筛查", command=self.start_scan)
        self.btn_scan.pack(side="left", padx=8, pady=8)
        self.btn_stop = ttk.Button(f3, text="停止", command=self.stop_scan, state="disabled")
        self.btn_stop.pack(side="left", padx=4)
        self.progress = ttk.Progressbar(f3, mode="determinate", length=420)
        self.progress.pack(side="left", padx=12)
        self.status = ttk.Label(f3, text="就绪")
        self.status.pack(side="left", padx=6)

        # 4. 结果
        f4 = ttk.LabelFrame(self.container, text="④ 命中结果（双击查看完整原文）")
        f4.pack(fill="both", expand=True, **pad)
        cols = ("sheet", "row", "person", "peer", "date", "kw", "cnt")
        self.tree = ttk.Treeview(f4, columns=cols, show="headings", selectmode="browse")
        heads = {"sheet": ("来源工作表", 150), "row": ("行号", 70), "person": ("行员", 100),
                 "peer": ("会话对象/客户", 180), "date": ("时间", 110), "kw": ("命中的关键词", 320), "cnt": ("命中次数", 80)}
        for c in cols:
            name, w = heads[c]
            self.tree.heading(c, text=name)
            self.tree.column(c, width=w, anchor="w")
        vsb = ttk.Scrollbar(f4, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(f4, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self.tree.grid(row=0, column=0, sticky="nsew", padx=(8, 0), pady=8)
        vsb.grid(row=0, column=1, sticky="ns", pady=8)
        hsb.grid(row=1, column=0, sticky="ew", padx=(8, 0))
        f4.rowconfigure(0, weight=1)
        f4.columnconfigure(0, weight=1)
        self.tree.bind("<Double-1>", lambda e: self.show_detail())

        # 5. 导出
        f5 = ttk.Frame(self.container)
        f5.pack(fill="x", **pad)
        self.summary = ttk.Label(f5, text="尚未筛查", foreground="#666666")
        self.summary.pack(side="left", padx=8)
        self.btn_stat = ttk.Button(f5, text="命中统计", command=self.show_stats, state="disabled")
        self.btn_stat.pack(side="right", padx=6)
        self.btn_export2 = ttk.Button(f5, text="仅导出命中行", command=lambda: self.do_export(only_hits=True), state="disabled")
        self.btn_export2.pack(side="right", padx=6)
        self.btn_export = ttk.Button(f5, text="导出标记版 Excel", command=lambda: self.do_export(False), state="disabled")
        self.btn_export.pack(side="right", padx=6)

        # 布局完成后再定尺寸，避免 macOS 上出现“有标题但内容空白”的窗口
        self.root.minsize(900, 640)
        self.root.geometry("1060x760")
        self.root.update_idletasks()

    # ---------------- 文件 ----------------
    def choose_file(self):
        path = filedialog.askopenfilename(
            title="选择聊天记录 Excel",
            filetypes=[("Excel 文件", "*.xlsx *.xlsm"), ("所有文件", "*.*")],
        )
        if path:
            self.set_file(path)

    def set_file(self, path):
        self.file_path = path
        self.file_var.set(path)
        self.file_info.configure(text="识别中…", foreground="#666666")
        self.root.update_idletasks()
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            names = wb.sheetnames
            wb.close()
            self.file_info.configure(text="已载入，工作表：" + " / ".join(names), foreground="#1a7f37")
        except Exception as e:
            self.file_info.configure(text=f"读取失败：{e}", foreground="#c00000")

    # ---------------- 关键词 ----------------
    def _kw_raw(self):
        return self.kw_text.get("1.0", "end-1c")

    def _update_kw_count(self):
        n = len(parse_keywords(self._kw_raw()))
        self.kw_count.configure(text=f"{n} 个关键词")

    def import_txt(self):
        path = filedialog.askopenfilename(title="导入关键词 txt", filetypes=[("文本文件", "*.txt"), ("所有文件", "*.*")])
        if not path:
            return
        try:
            txt = Path(path).read_text(encoding="utf-8", errors="ignore")
        except UnicodeDecodeError:
            txt = Path(path).read_text(encoding="gbk", errors="ignore")
        self.kw_text.insert("end", ("\n" if self._kw_raw().strip() else "") + txt)
        self._update_kw_count()

    def clear_kw(self):
        self.kw_text.delete("1.0", "end")
        self._update_kw_count()

    # ---------------- 扫描 ----------------
    def start_scan(self):
        if self.scanning:
            return
        if not self.file_path or not os.path.exists(self.file_path):
            messagebox.showwarning("提示", "请先选择待检查的 Excel 文件。")
            return
        keywords = parse_keywords(self._kw_raw())
        if not keywords:
            messagebox.showwarning("提示", "请输入至少一个关键词。")
            return
        config.save_keywords(self._kw_raw())

        self.result = None
        self.scanning = True
        self.cancel_event.clear()
        self.btn_scan.configure(state="disabled")
        self.btn_stop.configure(state="normal")
        self.btn_export.configure(state="disabled")
        self.btn_export2.configure(state="disabled")
        self.btn_stat.configure(state="disabled")
        self.progress["value"] = 0
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.status.configure(text="正在读取文件…")
        self.t0 = time.time()

        ignore = not self.case_sensitive.get()
        thread = threading.Thread(
            target=self._scan_worker, args=(self.file_path, keywords, ignore), daemon=True
        )
        thread.start()
        self.root.after(100, self._poll)

    def _scan_worker(self, path, keywords, ignore):
        def progress(done, total, found):
            self.q.put(("progress", done, total, found))

        try:
            res = scan_file(path, keywords, ignore_case=ignore,
                            progress=progress, should_cancel=self.cancel_event.is_set)
            self.q.put(("done", res))
        except Exception as e:
            self.q.put(("error", f"{type(e).__name__}: {e}"))

    def _poll(self):
        try:
            while True:
                item = self.q.get_nowait()
                kind = item[0]
                if kind == "progress":
                    _, done, total, found = item
                    self.progress["value"] = (done / total * 100) if total else 0
                    self.status.configure(text=f"已扫描 {done}/{total} 行，命中 {found} 条")
                elif kind == "done":
                    self._on_done(item[1])
                    return
                elif kind == "error":
                    self._finish()
                    messagebox.showerror("筛查失败", item[1])
                    return
        except queue.Empty:
            pass
        if self.scanning:
            self.root.after(100, self._poll)

    def _on_done(self, res):
        self.result = res
        self._finish()
        for h in res.hits:
            self.tree.insert("", "end", values=(
                h.sheet, h.row, h.person, h.peer, h.date,
                "、".join(h.keywords), sum(h.counts.values()),
            ), tags=(str(id(h)),))
        cost = time.time() - self.t0
        note = f"命中 {len(res.hits)} 条 / 扫描 {res.scanned} 行，耗时 {cost:.1f} 秒"
        if res.cancelled:
            note += "（已手动停止）"
        if res.truncated_rows:
            note += f"；⚠ {res.truncated_rows} 条疑似被 Excel 截断（超 32767 字），超出部分未参与匹配"
        self.summary.configure(text=note)
        self.status.configure(text="完成")
        if res.hits:
            self.btn_export.configure(state="normal")
            self.btn_export2.configure(state="normal")
            self.btn_stat.configure(state="normal")

    def _finish(self):
        self.scanning = False
        self.btn_scan.configure(state="normal")
        self.btn_stop.configure(state="disabled")
        self.progress["value"] = 100 if self.result else 0

    def stop_scan(self):
        if self.scanning:
            self.cancel_event.set()
            self.status.configure(text="正在停止…")

    # ---------------- 结果交互 ----------------
    def _selected_hit(self) -> HitRecord:
        sel = self.tree.selection()
        if not sel:
            return None
        idx = self.tree.index(sel[0])
        return self.result.hits[idx]

    def show_detail(self):
        h = self._selected_hit()
        if not h:
            return
        win = tk.Toplevel(self.root)
        win.title(f"原文 - {h.sheet} 第 {h.row} 行")
        win.geometry("900x600")
        info = (f"工作表：{h.sheet}    行号：{h.row}    行员：{h.person}    "
                f"会话对象：{h.peer}    时间：{h.date}\n"
                f"命中关键词：{'、'.join(h.keywords)}    "
                f"命中次数：{sum(h.counts.values())}"
                + ("    ⚠ 该条内容疑似被 Excel 截断" if h.truncated else ""))
        ttk.Label(win, text=info, wraplength=860, justify="left").pack(anchor="w", padx=10, pady=8)

        frame = ttk.Frame(win)
        frame.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        text = tk.Text(frame, wrap="word")
        text.pack(side="left", fill="both", expand=True)
        sb = ttk.Scrollbar(frame, command=text.yview)
        sb.pack(side="left", fill="y")
        text.configure(yscrollcommand=sb.set)
        text.insert("1.0", h.text)
        text.tag_config("hit", background=BG_HIT, foreground="#C00000")
        text.configure(state="disabled")

        cs = self.case_sensitive.get()
        hay = h.text if cs else h.text.lower()
        for kw in h.keywords:
            needle = kw if cs else kw.lower()
            start = 0
            while True:
                pos = hay.find(needle, start)
                if pos < 0:
                    break
                text.tag_add("hit", f"1.0+{pos}c", f"1.0+{pos + len(kw)}c")
                start = pos + len(needle)

    def show_stats(self):
        if not self.result:
            return
        c = Counter()
        for h in self.result.hits:
            for kw, n in h.counts.items():
                c[kw] += n
        win = tk.Toplevel(self.root)
        win.title("命中统计（按关键词）")
        win.geometry("420x520")
        tv = ttk.Treeview(win, columns=("kw", "n"), show="headings")
        tv.heading("kw", text="关键词")
        tv.heading("n", text="命中次数")
        tv.column("kw", width=260)
        tv.column("n", width=100, anchor="center")
        for kw, n in c.most_common():
            tv.insert("", "end", values=(kw, n))
        tv.pack(fill="both", expand=True, padx=10, pady=10)

    # ---------------- 导出 ----------------
    def do_export(self, only_hits: bool):
        if not self.result or not self.result.hits:
            return
        src = Path(self.file_path)
        default = src.with_name(src.stem + ("_命中清单.xlsx" if only_hits else "_筛查结果.xlsx"))
        dest = filedialog.asksaveasfilename(
            title="保存结果", defaultextension=".xlsx",
            initialfile=default.name, initialdir=str(src.parent),
            filetypes=[("Excel 文件", "*.xlsx")],
        )
        if not dest:
            return
        self.status.configure(text="正在导出…")
        self.root.update_idletasks()
        try:
            ignore = not self.case_sensitive.get()
            if only_hits:
                n = export_hits_only(str(src), dest, self.result.hits, ignore_case=ignore)
                msg = f"已导出 {n} 条命中记录，消息内容里的命中关键词已标红。\n{dest}"
            else:
                n = export_marked(str(src), dest, self.result.hits, ignore_case=ignore)
                msg = (f"已生成标记版文件，新增「是否命中」「命中的关键词」两列，"
                       f"其中 {n} 行标记为命中，消息内容里的命中关键词已标红。\n{dest}")
            self.status.configure(text="导出完成")
            messagebox.showinfo("导出完成", msg)
            if messagebox.askyesno("打开", "是否打开所在文件夹？"):
                self._reveal(dest)
        except Exception as e:
            self.status.configure(text="导出失败")
            messagebox.showerror("导出失败", f"{type(e).__name__}: {e}")

    @staticmethod
    def _reveal(path):
        import subprocess
        folder = os.path.dirname(os.path.abspath(path))
        try:
            if os.name == "nt":
                os.startfile(folder)  # noqa: S606
            else:
                subprocess.Popen(["open", folder])
        except Exception:
            pass

    # ---------------- 关闭 ----------------
    def on_close(self):
        config.save_keywords(self._kw_raw())
        self.root.destroy()


def main():
    root = tk.Tk()
    try:
        nb = ttk.Notebook(root)
        nb.pack(fill="both", expand=True)
        tab_scan = ttk.Frame(nb)
        tab_merge = ttk.Frame(nb)
        nb.add(tab_scan, text="筛查")
        nb.add(tab_merge, text="合并")
        app = App(root, tab_scan)
        merge_tab = MergeTab(tab_merge, app=app, notebook=nb)
        merge_tab.pack(fill="both", expand=True)
    except Exception:
        import traceback
        traceback.print_exc()
        messagebox.showerror("启动失败", traceback.format_exc())
        return
    # macOS：把窗口置顶并聚焦，避免窗口已创建但内容不绘制
    try:
        root.deiconify()
        root.lift()
        root.attributes("-topmost", True)
        root.after(300, lambda: root.attributes("-topmost", False))
        root.focus_force()
    except Exception:
        pass
    root.mainloop()


if __name__ == "__main__":
    main()
