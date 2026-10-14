"""导出：在原表基础上新增「是否命中」「命中的关键词」两列，并把消息内容里命中的关键词标红，原文件不修改。"""
from typing import List

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font, PatternFill

from scanner import CONTENT_HEADERS, HitRecord, _norm, _find_col

HIT_COL = "是否命中"
KW_COL = "命中的关键词"

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
HIT_FILL = PatternFill("solid", fgColor="FFF2CC")
HIT_FONT = InlineFont(color="FFFF0000", b=True)  # 命中的关键词：红色加粗


def _content_col(headers: List[str], fallback_index: int = None, ws=None):
    idx = _find_col(headers, CONTENT_HEADERS)
    if idx is not None:
        return idx
    if fallback_index is not None and ws is not None:
        return fallback_index
    return None


def _last_used_col(headers: List[str], ws) -> int:
    """返回最后一列的 1-based 列号。"""
    last = 0
    for i, h in enumerate(headers, start=1):
        if h:
            last = i
    if last == 0:
        last = ws.max_column or 1
    return last


def highlight_text(text, keywords: List[str], ignore_case: bool = True):
    """把命中关键词切成富文本：命中处红色加粗，其余保持原样（单元格内部分文字着色）。"""
    if not isinstance(text, str) or not text or not keywords:
        return text

    needles, seen = [], set()
    for kw in keywords:
        if not kw:
            continue
        nd = kw.lower() if ignore_case else kw
        if nd not in seen:
            seen.add(nd)
            needles.append(nd)

    hay = text.lower() if ignore_case else text
    spans = []
    for nd in needles:
        start = 0
        while True:
            pos = hay.find(nd, start)
            if pos < 0:
                break
            spans.append((pos, pos + len(nd)))
            start = pos + len(nd)
    if not spans:
        return text

    spans.sort()
    merged = [list(spans[0])]
    for s, e in spans[1:]:  # 合并重叠区间（关键词互相包含时）
        if s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])

    parts = []
    cursor = 0
    for s, e in merged:
        if s > cursor:
            parts.append(text[cursor:s])
        parts.append(TextBlock(HIT_FONT, text[s:e]))
        cursor = e
    if cursor < len(text):
        parts.append(text[cursor:])
    return CellRichText(*parts)


def export_marked(
    src_path: str,
    dest_path: str,
    hits: List[HitRecord],
    ignore_case: bool = True,
    highlight: bool = True,
) -> int:
    """生成标记版 Excel：每个被扫描的工作表末尾追加 是否命中 / 命中的关键词 两列，
    并把消息内容里命中的关键词标红。"""
    hit_map = {}
    for h in hits:
        hit_map.setdefault(h.sheet, {})[h.row] = h

    wb = load_workbook(src_path)
    marked_rows = 0
    for ws in wb.worksheets:
        if ws.title not in hit_map:
            continue
        headers = [_norm(c.value) for c in ws[1]]
        start = _last_used_col(headers, ws) + 1

        c1 = ws.cell(row=1, column=start, value=HIT_COL)
        c2 = ws.cell(row=1, column=start + 1, value=KW_COL)
        for c in (c1, c2):
            c.font = Font(bold=True)
            c.fill = HEADER_FILL

        content_idx = _find_col(headers, CONTENT_HEADERS)
        sheet_hits = hit_map[ws.title]
        for row in range(2, ws.max_row + 1):
            # 空行不标记，避免给 2 万多行空白行写入“否”
            if content_idx is not None:
                v = ws.cell(row=row, column=content_idx + 1).value
                if v is None or str(v).strip() == "":
                    continue
            h = sheet_hits.get(row)
            if h:
                ws.cell(row=row, column=start, value="是").fill = HIT_FILL
                ws.cell(row=row, column=start + 1, value="、".join(h.keywords))
                if highlight and content_idx is not None:
                    c = ws.cell(row=row, column=content_idx + 1)
                    if isinstance(c.value, str):
                        c.value = highlight_text(c.value, h.keywords, ignore_case)
                marked_rows += 1
            else:
                ws.cell(row=row, column=start, value="否")
                ws.cell(row=row, column=start + 1, value="")

    wb.save(dest_path)
    return marked_rows


def export_hits_only(
    src_path: str,
    dest_path: str,
    hits: List[HitRecord],
    ignore_case: bool = True,
    highlight: bool = True,
) -> int:
    """仅导出命中行（保留原表列 + 新增两列），消息内容里的关键词同样标红，便于分发复核。"""
    by_sheet = {}
    for h in hits:
        by_sheet.setdefault(h.sheet, []).append(h)

    src = load_workbook(src_path, read_only=True, data_only=True)
    out = Workbook()
    out.remove(out.active)
    total = 0

    for ws in src.worksheets:
        if ws.title not in by_sheet:
            continue
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            continue
        header = list(header)
        row_map = {}
        wanted = {h.row for h in by_sheet[ws.title]}
        for r_idx, row in enumerate(rows_iter, start=2):
            if r_idx in wanted:
                row_map[r_idx] = list(row)
            elif wanted and r_idx > max(wanted):
                break

        sheet = out.create_sheet(title=ws.title[:31])
        for i, v in enumerate(header, start=1):
            c = sheet.cell(row=1, column=i, value=v)
            c.font = Font(bold=True)
            c.fill = HEADER_FILL
        start = len(header) + 1
        c1 = sheet.cell(row=1, column=start, value=HIT_COL)
        c2 = sheet.cell(row=1, column=start + 1, value=KW_COL)
        for c in (c1, c2):
            c.font = Font(bold=True)
            c.fill = HEADER_FILL

        content_idx = _find_col([_norm(v) for v in header], CONTENT_HEADERS)

        for r, h in enumerate(sorted(by_sheet[ws.title], key=lambda x: x.row), start=2):
            values = row_map.get(h.row, [])
            for i, v in enumerate(values, start=1):
                sheet.cell(row=r, column=i, value=v)
            if highlight and content_idx is not None and content_idx < len(values):
                v = values[content_idx]
                if isinstance(v, str):
                    sheet.cell(row=r, column=content_idx + 1).value = highlight_text(
                        v, h.keywords, ignore_case
                    )
            sheet.cell(row=r, column=start, value="是").fill = HIT_FILL
            sheet.cell(row=r, column=start + 1, value="、".join(h.keywords))
            total += 1

    src.close()
    out.save(dest_path)
    return total
