"""聊天记录合并：同一「姓名 + 对方姓名 + 日期」的多条消息合并成一条。

规则（表 1，列名识别不到时按 C/G/J/K 位置回退）：
- 分组键：C 姓名 + G 对方姓名|昵称 + J 消息时间的年月日
- K 消息内容：按原表顺序拼接，每条前面带 [YYYY-MM-DD HH:MM:SS] 时间戳
- J 消息时间：只保留年月日（2026-09-01）
- 其余列：取该组第一条的值
- 结果写入新工作表（默认「合并结果」），默认保存在原文件里；也可另存为新文件
"""
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Callable, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from scanner import _norm

NAME_CANDIDATES = ("姓名", "行员姓名", "员工姓名")
PEER_CANDIDATES = ("对方姓名|昵称", "对方姓名", "客户姓名", "会话对象")
TIME_CANDIDATES = ("消息时间", "聊天时间", "时间", "日期")
CONTENT_CANDIDATES = ("消息内容", "聊天记录")

# 找不到列名时的回退位置（0-based）：C 姓名 / G 对方姓名 / J 消息时间 / K 消息内容
FALLBACK_COLS = {"name": 2, "peer": 6, "time": 9, "content": 10}

DATE_RE = re.compile(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})")
EXCEL_CELL_LIMIT = 32767
TRUNCATE_THRESHOLD = 32700
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
DEFAULT_SHEET = "合并结果"


@dataclass
class MergeResult:
    dest: str = ""
    sheet: str = ""
    rows_before: int = 0
    rows_after: int = 0
    longest: int = 0
    columns: Dict[str, str] = field(default_factory=dict)
    overflow: List[Tuple[int, int]] = field(default_factory=list)  # [(输出行号, 字符数)]
    cancelled: bool = False

    @property
    def saved(self) -> int:
        if not self.rows_before:
            return 0
        return round((1 - self.rows_after / self.rows_before) * 100)

    @property
    def summary(self) -> str:
        if not self.rows_after:
            return "没有可合并的数据"
        return f"{self.rows_before} 行 → {self.rows_after} 行（减少 {self.saved}%）"


def _cell(row, i: Optional[int]):
    if i is None or i >= len(row):
        return None
    return row[i]


def _is_blank(row) -> bool:
    return row is None or all(v is None or str(v).strip() == "" for v in row)


def date_key(v) -> str:
    """取年月日：2026-09-01 07:32:13 -> 2026-09-01。"""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    s = _norm(v)
    m = DATE_RE.search(s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    return s[:10]


def time_text(v) -> str:
    """完整时间戳：2026-09-01 07:32:13。"""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    return _norm(v)


def detect_columns(headers: List[str]) -> Tuple[Dict[str, int], Dict[str, str]]:
    """按列名定位 4 个关键列，识别不到时用 C/G/J/K 回退。"""
    norm = [_norm(h) for h in headers]

    def find(candidates, fallback):
        for name in candidates:
            if name in norm:
                return norm.index(name), name
        i = fallback
        label = norm[i] if i < len(norm) and norm[i] else f"第 {get_column_letter(i + 1)} 列"
        return i, label

    idx, labels = {}, {}
    for key, cands in (
        ("name", NAME_CANDIDATES),
        ("peer", PEER_CANDIDATES),
        ("time", TIME_CANDIDATES),
        ("content", CONTENT_CANDIDATES),
    ):
        idx[key], labels[key] = find(cands, FALLBACK_COLS[key])
    return idx, labels


def _unique_title(wb, title: str) -> str:
    """避免覆盖已有工作表：重名则追加 _2 / _3 …"""
    name = (title or DEFAULT_SHEET).strip()[:31] or DEFAULT_SHEET
    if name not in wb.sheetnames:
        return name
    i = 2
    while True:
        suffix = f"_{i}"
        candidate = name[: 31 - len(suffix)] + suffix
        if candidate not in wb.sheetnames:
            return candidate
        i += 1


def inspect_sheet(path: str, sheet_index: int = 0) -> dict:
    """只读表头和行数，给界面做预览（比正式合并便宜很多）。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[sheet_index]
        header = [_norm(v) for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()) or ()]
        idx, labels = detect_columns(header)
        rows = max((ws.max_row or 1) - 1, 0)
        return {
            "sheet": ws.title,
            "sheets": wb.sheetnames,
            "rows": rows,
            "columns": labels,
            "headers": header,
            "index": idx,
        }
    finally:
        wb.close()


def merge_sheet(
    path: str,
    dest_path: Optional[str] = None,
    sheet_title: str = DEFAULT_SHEET,
    sheet_index: int = 0,
    progress: Optional[Callable[[int, int, str], None]] = None,
    should_cancel: Optional[Callable[[], bool]] = None,
) -> MergeResult:
    """合并指定文件的第 sheet_index 个工作表，结果写入新工作表。"""
    result = MergeResult()
    wb = load_workbook(path)
    try:
        ws = wb.worksheets[sheet_index]
        header = [_norm(v) for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()) or ()]
        if not any(header):
            raise ValueError("该工作表没有表头，无法识别列。")
        idx, labels = detect_columns(header)
        result.columns = labels

        # ---- 1. 读：按 (姓名, 对方姓名, 日期) 分组 ----
        total = max((ws.max_row or 1) - 1, 0)
        groups, order = {}, []
        done = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if _is_blank(row):
                continue
            done += 1
            key = (
                _norm(_cell(row, idx["name"])),
                _norm(_cell(row, idx["peer"])),
                date_key(_cell(row, idx["time"])),
            )
            g = groups.get(key)
            if g is None:
                g = {"first": row, "parts": []}
                groups[key] = g
                order.append(key)
            text = _norm(_cell(row, idx["content"]))
            ts = time_text(_cell(row, idx["time"]))
            g["parts"].append(f"[{ts}]\n{text}" if text else f"[{ts}]")
            if progress and done % 500 == 0:
                progress(done, total + len(groups) + 1, "读取原表")
            if should_cancel and should_cancel():
                result.cancelled = True
                break

        result.rows_before = done
        if result.cancelled or not groups:
            return result

        # ---- 2. 写：每组一行 ----
        out_title = _unique_title(wb, sheet_title)
        out = wb.create_sheet(out_title)
        for i, v in enumerate(header, start=1):
            c = out.cell(row=1, column=i, value=v)
            c.font = Font(bold=True)
            c.fill = HEADER_FILL

        written = 0
        for key in order:
            g = groups[key]
            first = g["first"]
            content = "\n".join(g["parts"])
            if len(content) > result.longest:
                result.longest = len(content)
            for i in range(len(header)):
                if i == idx["time"]:
                    v = date_key(_cell(first, i))  # 只留年月日
                elif i == idx["content"]:
                    v = content
                else:
                    v = _cell(first, i)  # 其余列取第一条
                out.cell(row=written + 2, column=i + 1, value=v)
            if len(content) >= TRUNCATE_THRESHOLD:
                result.overflow.append((written + 2, len(content)))
            written += 1
            if progress and written % 200 == 0:
                progress(done + written, total + len(order) + 1, "写入结果")
            if should_cancel and should_cancel():
                result.cancelled = True
                break

        result.rows_after = written
        result.sheet = out_title

        # ---- 3. 保存 ----
        dest = dest_path or path
        if progress:
            progress(done + written, total + len(order) + 1, "保存文件")
        wb.save(dest)
        result.dest = dest
    finally:
        wb.close()
    return result
